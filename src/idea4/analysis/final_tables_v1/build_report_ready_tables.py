from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT_DIR = Path(__file__).resolve().parent

MAIN_IN = OUT_DIR / "main_k_scaling_summary.csv"
STRESS_IN = OUT_DIR / "stress_complexity_summary.csv"
ABLATION_IN = OUT_DIR / "ablation_summary.csv"

MAIN_REPORT_CSV = OUT_DIR / "report_main_k_scaling.csv"
STRESS_REPORT_CSV = OUT_DIR / "report_stress_complexity.csv"
ABLATION_REPORT_CSV = OUT_DIR / "report_ablation.csv"
REPORT_WORKBOOK = OUT_DIR / "report_ready_tables.xlsx"

MAIN_MODEL_ORDER = {
    "DQN baseline": 1,
    "Basic PPO": 2,
    "Factorized AC": 3,
    "Dual-branch AC": 4,
}

STRESS_MODEL_ORDER = {
    "DQN baseline": 1,
    "Basic PPO": 2,
    "Factorized AC": 3,
}

MAIN_NOTE = (
    "This table compares the main models under C=1200, F=3, T=1000 as the number of wallets k increases. "
    "Confidence intervals are computed across random seeds. The main pattern is that flat DQN degrades strongly "
    "as k increases, while factorized actor-critic models maintain substantially higher acceptance at larger k."
)

STRESS_NOTE = (
    "This table evaluates harder settings, including larger action spaces and tighter capacity. Confidence intervals "
    "are computed across random seeds. The results show that factorized policies generally remain more competitive "
    "than flat DQN under stress, although performance becomes unstable in the hardest C=800, k=24 setting."
)

ABLATION_NOTE = (
    "This table summarizes exploratory ablations. Some variants only have one seed, so their confidence intervals "
    "are reported as NA. The gate-balanced dual actor-critic variant is the strongest ablation at C=1200, k=12, F=3, "
    "where it improves mean value acceptance over the base dual-branch model."
)

README_LINES = [
    "These tables are presentation-ready summaries derived from final_tables_v1.",
    "CI is computed across random seeds using a t-based 95% confidence interval.",
    "If only one seed is available, CI is shown as NA.",
    "The best model in each comparable group is bolded based on highest Mean Value Acceptance.",
    "Original reward results only; money reward experiments are not included.",
]


def fmt_number(value: Any, decimals: int = 2) -> str:
    if pd.isna(value):
        return "NA"
    return f"{float(value):.{decimals}f}"


def fmt_ci(mean: Any, ci: Any, percent: bool = False) -> str:
    if pd.isna(mean):
        mean_text = "NA"
    elif percent:
        mean_text = f"{float(mean) * 100:.2f}%"
    else:
        mean_text = fmt_number(mean)

    if pd.isna(ci):
        ci_text = "NA"
    elif percent:
        ci_text = f"{float(ci) * 100:.2f}%"
    else:
        ci_text = fmt_number(ci)
    return f"{mean_text} +/- {ci_text}"


def add_best_flags(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    df = df.copy()
    df["_best_mean_value_acceptance"] = False
    df["_lowest_drops"] = False
    for _, group in df.groupby(group_cols, dropna=False):
        if group["mean_value_accept_ratio_mean"].notna().any():
            best_idx = group["mean_value_accept_ratio_mean"].idxmax()
            df.loc[best_idx, "_best_mean_value_acceptance"] = True
        if group["drops_mean"].notna().any():
            low_idx = group["drops_mean"].idxmin()
            df.loc[low_idx, "_lowest_drops"] = True
    return df


def build_main_report() -> pd.DataFrame:
    df = pd.read_csv(MAIN_IN)
    df["_model_order"] = df["model_clean"].map(MAIN_MODEL_ORDER).fillna(99)
    df = df.sort_values(["k", "_model_order"]).copy()
    df = add_best_flags(df, ["k"])
    return pd.DataFrame(
        {
            "k": df["k"].astype(int),
            "Model": df["model_clean"],
            "Seeds": df["n_seeds"].astype(int),
            "Mean Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["mean_value_accept_ratio_mean"], df["mean_value_accept_ratio_ci95"])
            ],
            "Worst-Regime Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["worst_regime_value_accept_ratio_mean"], df["worst_regime_value_accept_ratio_ci95"])
            ],
            "Drops +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["drops_mean"], df["drops_ci95"])],
            "Flushes +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["flushes_mean"], df["flushes_ci95"])],
            "_best_mean_value_acceptance": df["_best_mean_value_acceptance"],
            "_lowest_drops": df["_lowest_drops"],
        }
    )


def build_stress_report() -> pd.DataFrame:
    df = pd.read_csv(STRESS_IN)
    df["_model_order"] = df["model_clean"].map(STRESS_MODEL_ORDER).fillna(99)
    df = df.sort_values(["C", "k", "_model_order"]).copy()
    df = add_best_flags(df, ["C", "k"])
    return pd.DataFrame(
        {
            "C": df["C"].astype(int),
            "k": df["k"].astype(int),
            "Model": df["model_clean"],
            "Seeds": df["n_seeds"].astype(int),
            "Mean Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["mean_value_accept_ratio_mean"], df["mean_value_accept_ratio_ci95"])
            ],
            "Worst-Regime Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["worst_regime_value_accept_ratio_mean"], df["worst_regime_value_accept_ratio_ci95"])
            ],
            "Drops +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["drops_mean"], df["drops_ci95"])],
            "Flushes +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["flushes_mean"], df["flushes_ci95"])],
            "_best_mean_value_acceptance": df["_best_mean_value_acceptance"],
            "_lowest_drops": df["_lowest_drops"],
        }
    )


def build_ablation_report() -> pd.DataFrame:
    df = pd.read_csv(ABLATION_IN)
    df = df.sort_values(["C", "k", "F", "model_clean", "variant_clean"]).copy()
    df = add_best_flags(df, ["C", "k", "F"])
    return pd.DataFrame(
        {
            "C": df["C"].astype(int),
            "k": df["k"].astype(int),
            "F": df["F"].astype(int),
            "Model": df["model_clean"],
            "Variant": df["variant_clean"],
            "Seeds": df["n_seeds"].astype(int),
            "Mean Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["mean_value_accept_ratio_mean"], df["mean_value_accept_ratio_ci95"])
            ],
            "Worst-Regime Value Acceptance +/- 95% CI": [
                fmt_ci(m, c, percent=True)
                for m, c in zip(df["worst_regime_value_accept_ratio_mean"], df["worst_regime_value_accept_ratio_ci95"])
            ],
            "Drops +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["drops_mean"], df["drops_ci95"])],
            "Flushes +/- 95% CI": [fmt_ci(m, c) for m, c in zip(df["flushes_mean"], df["flushes_ci95"])],
            "CI Note": df["ci_note"],
            "_best_mean_value_acceptance": df["_best_mean_value_acceptance"],
            "_lowest_drops": df["_lowest_drops"],
        }
    )


def visible(df: pd.DataFrame) -> pd.DataFrame:
    return df[[col for col in df.columns if not col.startswith("_")]]


def write_csvs(main: pd.DataFrame, stress: pd.DataFrame, ablation: pd.DataFrame) -> None:
    visible(main).to_csv(MAIN_REPORT_CSV, index=False)
    visible(stress).to_csv(STRESS_REPORT_CSV, index=False)
    visible(ablation).to_csv(ABLATION_REPORT_CSV, index=False)


def write_workbook(main: pd.DataFrame, stress: pd.DataFrame, ablation: pd.DataFrame) -> None:
    with pd.ExcelWriter(REPORT_WORKBOOK, engine="openpyxl") as writer:
        pd.DataFrame({"README": README_LINES}).to_excel(writer, sheet_name="README", index=False)
        visible(main).to_excel(writer, sheet_name="Main K Scaling Report", index=False)
        visible(stress).to_excel(writer, sheet_name="Stress Complexity Report", index=False)
        visible(ablation).to_excel(writer, sheet_name="Ablation Report", index=False)

    wb = load_workbook(REPORT_WORKBOOK)
    format_readme(wb["README"])
    format_report_sheet(wb["Main K Scaling Report"], main, MAIN_NOTE)
    format_report_sheet(wb["Stress Complexity Report"], stress, STRESS_NOTE)
    format_report_sheet(wb["Ablation Report"], ablation, ABLATION_NOTE)
    wb.save(REPORT_WORKBOOK)


def format_readme(ws) -> None:
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 110
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def format_report_sheet(ws, full_df: pd.DataFrame, note: str) -> None:
    visible_cols = [col for col in full_df.columns if not col.startswith("_")]
    max_row = len(full_df) + 1
    max_col = len(visible_cols)

    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    best_fill = PatternFill("solid", fgColor="E2F0D9")
    low_drop_fill = PatternFill("solid", fgColor="FFF2CC")
    thin_side = Side(style="thin", color="D9D9D9")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(2, max_row + 1):
        is_best = bool(full_df.iloc[row_idx - 2]["_best_mean_value_acceptance"])
        is_low_drop = bool(full_df.iloc[row_idx - 2]["_lowest_drops"])
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if is_best:
                cell.font = Font(bold=True)
                cell.fill = best_fill
            elif is_low_drop:
                cell.fill = low_drop_fill

    widths = {
        "C": 10,
        "k": 8,
        "F": 8,
        "Model": 26,
        "Variant": 20,
        "Seeds": 10,
        "Mean Value Acceptance +/- 95% CI": 26,
        "Worst-Regime Value Acceptance +/- 95% CI": 30,
        "Drops +/- 95% CI": 20,
        "Flushes +/- 95% CI": 20,
        "CI Note": 34,
    }
    for idx, col_name in enumerate(visible_cols, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col_name, 18)

    note_row = max_row + 2
    ws.cell(row=note_row, column=1, value="Note").font = Font(bold=True)
    ws.cell(row=note_row + 1, column=1, value=note)
    ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=max_col)
    ws.cell(row=note_row + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=note_row + 1, column=1).fill = PatternFill("solid", fgColor="F2F2F2")


def main() -> None:
    main_report = build_main_report()
    stress_report = build_stress_report()
    ablation_report = build_ablation_report()
    write_csvs(main_report, stress_report, ablation_report)
    write_workbook(main_report, stress_report, ablation_report)

    print(f"Wrote {MAIN_REPORT_CSV}")
    print(f"Wrote {STRESS_REPORT_CSV}")
    print(f"Wrote {ABLATION_REPORT_CSV}")
    print(f"Wrote {REPORT_WORKBOOK}")


if __name__ == "__main__":
    main()
